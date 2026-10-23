"""Run the whole QA suite (database, API, database -> API trace, UI) and write the QA workbook.

    python tests/qa/run_qa.py                 everything, local sites (servers :8600 :8601 :5178 :5180 must be running)
    python tests/qa/run_qa.py --no-ui         database + API only
    python tests/qa/run_qa.py --report-only   rebuild the workbook from the last results in tests/qa/out/

Other targets (e.g. the VPS, read-only): set KSSL_QA_* (see conftest.py and playwright.qa.config.mjs).
Writes tests/qa/out/KSSL_QA_report.xlsx: Summary, Test cases (one row per case: steps, expected, result), Results (every
check), Findings (each failing case with its triage from FINDINGS below).
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = HERE / "out"

# Triage of the failures the first run found (2026-10-23): severity, what it means, what to do. A failing case not listed
# here shows as "to triage" in the Findings sheet.
FINDINGS = {
    "DB-006": ("Medium", "149 of 186 card / ui tables have no PRIMARY KEY; uniqueness rests on the loaders alone (DB-011 checks it in the data).",
               "Add PRIMARY KEY on each card model key column (backend/review.py GRAPH) in a migration, staging and serving."),
    "DB-012": ("Medium", "Competitor e7b7d2e6-cd75-5963-9669-1c92c49b7969 was removed but its locations (2), leaders (4) and news articles (34) remain, in staging and serving.",
               "Delete the orphan rows (or restore the competitor); add FOREIGN KEYs / a cascade in the removal script."),
    "DB-021": ("Low", "Some source urls are not web addresses: 1,433 fact_source rows for the image_local_path column hold the stored image path ('news/<hash>.jpg'); a L&T PDF url has spaces and is cut short ('...BUSINESS.pd').",
               "Record the image's web url (media.image.source_url) as its source; repair the truncated PDF url."),
    "DB-041": ("High", "Sources point at row ids that no longer exist: derived tables (matchups, matchup advantages, top opportunities, innovation status/company/milestones) were rebuilt with new ids and their fact_source rows were not moved -- the new rows show no source.",
               "Re-key fact_source when these tables are rebuilt (or use stable ids); then DB-050 drops for top opportunities / matchups."),
    "DB-050": ("High", "Review cards without any source in their bundle: matchups (judgement), country list (geo_location), top-opportunity / next-move rankings, stories with no sourced step, market / patent field assessments, 1 partner, 1 partnership, 1 field synthesis.",
               "Owner rule: a source on everything in review. Derived cards should cite what they derive from (tender / product / signal sources); fix DB-041 first."),
    "DB-051": ("High", "A live news-derived card with no source (1 partnership).", "Find its article or unpublish it."),
    "DB-052": ("Medium", "Incomplete source records: 39 with no source_tier; 30 whose url is a list fragment (',' or ',https://...'); plus the 1,433 image-path rows of DB-021.",
               "Fix the writers (tier always set, one url per row); clean the rows."),
    "UI-011": ("Low", "Search cannot find a title typed (or pasted) with its original en dash: the page rewrites dashes to hyphens (src/lib/dataset.js) but the query is not.",
               "Apply the same dash normalisation to the query in src/lib/search.js."),
    "API-041": ("Known", "The temporary login is a UI gate only: the API answers without a login.", "The real login must lock the API (planned)."),
}


def run(cmd: list[str]) -> int:
    print("$", " ".join(cmd), flush=True)
    return subprocess.call(cmd, cwd=ROOT)


def ui_results() -> list[dict]:
    p = OUT / "ui_results.json"
    if not p.exists():
        return []
    out = []

    def walk(suite):
        for s in suite.get("suites", []):
            walk(s)
        for spec in suite.get("specs", []):
            for t in spec.get("tests", []):
                ann = {a["type"]: a.get("description", "") for r in t.get("results", [])[-1:] for a in r.get("annotations", [])}
                ann = ann or {a["type"]: a.get("description", "") for a in t.get("annotations", [])}
                last = (t.get("results") or [{}])[-1]
                status = {"expected": "passed", "unexpected": "failed", "flaky": "passed", "skipped": "skipped"}.get(t.get("status"), t.get("status"))
                err = re.sub(r"\x1b\[[0-9;]*m", "", "\n".join((e.get("message") or "")[:600] for e in last.get("errors", [])))
                tid, _, title = spec["title"].partition("·")
                case = tid.strip()
                out.append({"id": f"{case}[{t.get('projectName')}]", "case": case, "title": title.strip(), "param": t.get("projectName"),
                            "area": "UI", "steps": " ".join(ann.get("steps", "").split()), "expected": " ".join(ann.get("expected", "").split()),
                            "outcome": status, "detail": err, "seconds": round(last.get("duration", 0) / 1000, 2),
                            "flaky": t.get("status") == "flaky"})

    for s in json.loads(p.read_text(encoding="utf-8")).get("suites", []):
        walk(s)
    return out


def workbook(results: list[dict]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3A5F")
    colour = {"passed": "D9EAD3", "failed": "F4CCCC", "skipped": "EEEEEE", "xfailed": "FFF2CC", "error": "F4CCCC"}

    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def sheet(ws, header, rows_, widths):
        rows_ = [[ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in r] for r in rows_]
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows_:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    layer = {"DB": "1 Database", "API": "2 Backend API", "TRACE": "3 Database -> API", "UI": "4 UI (database -> screen)"}
    cases: dict[str, dict] = {}
    for r in results:
        c = cases.setdefault(r["case"], {**r, "n": 0, "passed": 0, "failed": 0, "skipped": 0, "fails": []})
        c["n"] += 1
        k = "failed" if r["outcome"] in ("failed", "error") else "skipped" if r["outcome"] in ("skipped", "xfailed") else "passed"
        c[k] += 1
        if k == "failed":
            c["fails"].append(f"[{r['param'] or '-'}] {r['detail'].splitlines()[0] if r['detail'] else ''}"[:400])

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["KSSL QA report", dt.datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Targets", os.environ.get("KSSL_QA_TARGET", "local: DBs kssl_staging / kssl_serving, APIs :8600 / :8601, sites :5178 / :5180")])
    ws.append([])
    ws.append(["Layer", "Test cases", "Checks run", "Passed", "Failed", "Skipped / known"])
    for c in ws[4]:
        c.font, c.fill = head, fill
    tot = [0, 0, 0, 0, 0]
    for key, name in layer.items():
        cs = [c for c in cases.values() if c["area"] == key]
        rs = [r for r in results if r["area"] == key]
        p = sum(1 for r in rs if r["outcome"] == "passed")
        f = sum(1 for r in rs if r["outcome"] in ("failed", "error"))
        s = len(rs) - p - f
        ws.append([name, len(cs), len(rs), p, f, s])
        tot = [a + b for a, b in zip(tot, [len(cs), len(rs), p, f, s])]
    ws.append(["Total", *tot])
    ws.append([])
    ws.append(["Failing test cases", sum(1 for c in cases.values() if c["failed"])])
    ws.append(["How to re-run", "python tests/qa/run_qa.py   (see tests/qa/README.md)"])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 90
    for col in "CDEF":
        ws.column_dimensions[col].width = 14

    order = lambda c: (layer.get(c["area"], c["area"]), c["case"])  # noqa: E731
    sheet(wb.create_sheet("Test cases"),
          ["ID", "Layer", "Test case", "Steps", "Expected result", "Checks", "Passed", "Failed", "Skipped", "Status", "Failures (first 5)"],
          [[c["case"], layer.get(c["area"], c["area"]), c["title"], c["steps"], c["expected"], c["n"], c["passed"], c["failed"], c["skipped"],
            "FAIL" if c["failed"] else ("PASS" if c["passed"] else "SKIP"), "\n".join(c["fails"][:5])] for c in sorted(cases.values(), key=order)],
          [10, 16, 40, 55, 55, 8, 8, 8, 8, 8, 70])
    ws = wb["Test cases"]
    for row in ws.iter_rows(min_row=2):
        row[9].fill = PatternFill("solid", fgColor={"FAIL": "F4CCCC", "PASS": "D9EAD3"}.get(row[9].value, "EEEEEE"))

    sheet(wb.create_sheet("Results"), ["Check", "Layer", "Test case", "Parameter", "Result", "Seconds", "Detail"],
          [[r["id"], layer.get(r["area"], r["area"]), r["title"], r["param"], r["outcome"], r["seconds"], r["detail"][:2000]]
           for r in sorted(results, key=lambda r: (layer.get(r["area"], ""), r["id"]))], [45, 16, 40, 30, 10, 9, 90])
    ws = wb["Results"]
    for row in ws.iter_rows(min_row=2):
        row[4].fill = PatternFill("solid", fgColor=colour.get(row[4].value, "FFFFFF"))

    fails = [c for c in sorted(cases.values(), key=order) if c["failed"] or c["case"] in FINDINGS and c["skipped"]]
    sheet(wb.create_sheet("Findings"), ["ID", "Test case", "Severity", "What it means", "Recommended fix", "Failing checks", "Examples"],
          [[c["case"], c["title"], *(FINDINGS.get(c["case"], ("to triage", "", ""))), c["failed"] or c["skipped"], "\n".join(c["fails"][:5])]
           for c in fails], [10, 38, 10, 60, 50, 9, 70])
    path = OUT / "KSSL_QA_report.xlsx"
    wb.save(path)
    return path


def main() -> int:
    OUT.mkdir(exist_ok=True)
    if "--report-only" not in sys.argv:
        run([sys.executable, "-m", "pytest", "tests/qa", "-q", "-p", "no:cacheprovider", "-W", "ignore"])
        if "--no-ui" not in sys.argv:
            run(["npx", "playwright", "test", "-c", "playwright.qa.config.mjs"] if os.name != "nt" else
                ["cmd", "/c", "npx", "playwright", "test", "-c", "playwright.qa.config.mjs"])
    py = json.loads((OUT / "pytest_results.json").read_text(encoding="utf-8")) if (OUT / "pytest_results.json").exists() else []
    results = py + ui_results()
    path = workbook(results)
    bad = sum(1 for r in results if r["outcome"] in ("failed", "error"))
    print(f"{len(results)} checks, {bad} failed -> {path}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
